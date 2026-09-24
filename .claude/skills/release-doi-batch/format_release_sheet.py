#!/usr/bin/env python3
"""
Format a per-release curation tracking spreadsheet into the DOIs.xlsx worksheet
layout used by previous releases (V95, V96, ...), ready for
generate_crossref_xml.py.

The curation team tracks a release in a wide working sheet (e.g. "Release 97.xlsx")
with columns like Curator, External_Contributor, contributor_type, ORCID,
contributor_URL, TLP, release_site, NOTES ... . generate_crossref_xml.py expects
the narrower, canonical DOIs.xlsx layout:

    Project_ID, Project, Curator Given, Curator Surname, Curator ORCID,
    StableID, StableID_root, Contributor Given, Contributor Surname,
    Contributor Role, Contributor ORCID, Editor, Reviewer, Title, DOI Data

This script performs that conversion, validates the result, and either writes a
standalone workbook or inserts the sheet into DOIs.xlsx in place (with a backup).

Usage:
    python3 format_release_sheet.py V97 --source "Release 97.xlsx" \
        --dois /path/to/DOIs.xlsx --output ./output/V97-doi-batch/V97_formatted.xlsx

    # validate only, write nothing
    python3 format_release_sheet.py V97 --source "Release 97.xlsx" \
        --dois /path/to/DOIs.xlsx --validate-only

    # insert the formatted sheet straight into DOIs.xlsx (backs the file up first)
    python3 format_release_sheet.py V97 --source "Release 97.xlsx" \
        --dois /path/to/DOIs.xlsx --in-place
"""

import argparse
import re
import shutil
import sys
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter


# Canonical DOIs.xlsx worksheet layout, in order. Do not reorder — downstream
# consumers and prior releases rely on these exact header strings.
TARGET_COLUMNS = [
    "Project_ID",
    "Project",
    "Curator Given",
    "Curator Surname",
    "Curator ORCID",
    "StableID",
    "StableID_root",
    "Contributor Given",
    "Contributor Surname",
    "Contributor Role",
    "Contributor ORCID",
    "Editor",
    "Reviewer",
    "Title",
    "DOI Data",
]

# Source column aliases, normalized (lowercased, non-alphanumerics -> "_").
SOURCE_ALIASES = {
    "project_id": ["project_id", "projectid", "db_id", "dbid"],
    "project": ["project", "pathway", "pathway_name", "name"],
    "curator": ["curator", "curator_name", "curator_full_name"],
    "stable_id": ["stableid", "stable_id"],
    "stable_id_root": ["stableid_root", "stable_id_root", "stableidroot"],
    "contributor": ["external_contributor", "contributor", "contributor_name",
                    "contributor_full_name"],
    "contributor_last": ["contributor_last_name", "contributor_surname",
                         "contributor_lastname"],
    "role": ["contributor_type", "contributor_role", "role"],
    "orcid": ["orcid", "contributor_orcid", "orcid_id"],
    "contributor_url": ["contributor_url", "orcid_url", "contributor_orcid_url"],
    "declared_doi": ["new_updated_doi", "new_doi", "doi", "updated_doi"],
    "no_orcid": ["no_orcid_id", "no_orcid"],
}

STABLE_ID_RE = re.compile(r"^R-[A-Z]{3}-\d+\.\d+$")
ORCID_RE = re.compile(r"^\d{4}-\d{4}-\d{4}-\d{3}[\dX]$")
DOI_PREFIX = "10.3180"
RESOURCE_BASE = "http://reactome.org/content/detail"


class Report:
    """Collects validation findings, separated into blocking and advisory."""

    def __init__(self):
        self.errors = []
        self.warnings = []

    def error(self, row, message):
        self.errors.append((row, message))

    def warn(self, row, message):
        self.warnings.append((row, message))

    def print_report(self):
        def dump(label, items):
            if not items:
                return
            print(f"\n{label} ({len(items)}):")
            for row, message in items:
                where = f"source row {row}" if row is not None else "workbook"
                print(f"  - [{where}] {message}")

        dump("ERRORS", self.errors)
        dump("WARNINGS", self.warnings)
        if not self.errors and not self.warnings:
            print("\nValidation: clean — no errors, no warnings.")


def normalize(name):
    return re.sub(r"[^a-z0-9]+", "_", str(name).strip().lower()).strip("_")


def clean(value):
    """Collapse whitespace/newlines in a cell value; NaN and blanks become ''."""
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return ""
    if pd.isna(value):
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def resolve_columns(df, report):
    """Map normalized source headers onto our logical field names."""
    lookup = {}
    for col in df.columns:
        lookup.setdefault(normalize(col), col)

    resolved = {}
    for field, aliases in SOURCE_ALIASES.items():
        for alias in aliases:
            if alias in lookup:
                resolved[field] = lookup[alias]
                break

    required = ["project_id", "project", "curator", "stable_id", "contributor"]
    missing = [f for f in required if f not in resolved]
    if missing:
        sys.exit(
            "Error: source sheet is missing required column(s): "
            + ", ".join(missing)
            + f"\n  Columns found: {list(df.columns)}"
        )
    return resolved


def load_curator_orcids(dois_path, report):
    """Build {'given surname': orcid} from the 'Curator ORCID' sheet."""
    try:
        sheet = pd.read_excel(dois_path, sheet_name="Curator ORCID")
    except ValueError:
        report.warn(None, "DOIs.xlsx has no 'Curator ORCID' sheet — curator "
                          "ORCIDs will be left blank.")
        return {}

    sheet.columns = [str(c).strip() for c in sheet.columns]
    table = {}
    for _, row in sheet.iterrows():
        given = clean(row.get("Curator Given"))
        surname = clean(row.get("Curator Surname"))
        orcid = clean(row.get("ORCID"))
        if given or surname:
            table[f"{given} {surname}".strip().lower()] = (given, surname, orcid)
    return table


def split_curator(full_name, curator_table, row_num, report):
    """Split a curator's full name, preferring an exact Curator ORCID match."""
    name = clean(full_name)
    if not name:
        report.error(row_num, "curator name is blank")
        return "", "", ""

    hit = curator_table.get(name.lower())
    if hit:
        return hit

    # Fall back to a last-space split and flag it — the curator is not in the
    # lookup sheet, so their ORCID cannot be filled in.
    parts = name.split()
    given, surname = (" ".join(parts[:-1]), parts[-1]) if len(parts) > 1 else ("", name)
    report.error(
        row_num,
        f"curator {name!r} is not in the 'Curator ORCID' sheet — split as "
        f"given={given!r} surname={surname!r} with no ORCID. Add them to the "
        f"lookup sheet.",
    )
    return given, surname, ""


def split_contributor(full_name, last_name, row_num, report):
    """Split a contributor's full name using the supplied surname as the anchor.

    Anchoring on the surname column keeps multi-word given names ("David P")
    and hyphenated/compound surnames ("Senff-Ribeiro") intact.
    """
    name = clean(full_name)
    surname = clean(last_name)
    if not name:
        return "", ""

    if surname and name.lower().endswith(surname.lower()):
        given = name[: len(name) - len(surname)].strip()
        return given, surname

    if surname:
        report.warn(
            row_num,
            f"contributor {name!r} does not end with the stated surname "
            f"{surname!r} — splitting on the last space instead.",
        )
    parts = name.split()
    if len(parts) > 1:
        return " ".join(parts[:-1]), parts[-1]
    return "", name


def resolve_orcid(orcid_value, url_value, row_num, who, report):
    """Take the ORCID column, else recover one from the contributor URL."""
    orcid = clean(orcid_value)
    if not orcid or orcid.upper() in {"N/A", "NA", "NONE", "NAN"}:
        url = clean(url_value)
        match = re.search(r"(\d{4}-\d{4}-\d{4}-\d{3}[\dX])", url, re.IGNORECASE)
        orcid = match.group(1).upper() if match else ""

    if not orcid:
        report.warn(row_num, f"{who} has no ORCID — the <ORCID> element will be "
                             f"omitted for this contributor.")
        return ""

    if not ORCID_RE.match(orcid):
        report.error(row_num, f"{who} ORCID {orcid!r} is not a valid ORCID "
                              f"(expected 0000-0000-0000-0000).")
    return orcid


def person_xml(given, surname, orcid):
    """Legacy helper-column fragment, matching the V95/V96 Editor/Reviewer cells.

    Note: where a contributor has no ORCID, older sheets emitted an empty
    'https://orcid.org/' value. We omit the <ORCID> line instead, which matches
    what generate_crossref_xml.py actually deposits.
    """
    lines = [f"<given_name>{given}</given_name>", f"<surname>{surname}</surname>"]
    if orcid:
        lines.append(f'<ORCID authenticated="true">https://orcid.org/{orcid}</ORCID>')
    lines.append("</person_name>")
    return "\n".join(lines)


def build_rows(src, resolved, curator_table, report):
    """Convert each source row into a canonical DOIs.xlsx row."""
    rows = []
    for idx, raw in src.iterrows():
        row_num = idx + 2  # 1-based with a header row, matching what Excel shows

        def field(name):
            col = resolved.get(name)
            return raw[col] if col is not None else None

        project = clean(field("project"))
        stable_id = clean(field("stable_id"))
        project_id_raw = clean(field("project_id"))

        # Skip padding rows left at the bottom of the tracking sheet.
        if not project and not stable_id and not project_id_raw:
            continue

        # --- Project_ID -----------------------------------------------------
        try:
            project_id = int(float(project_id_raw))
        except (TypeError, ValueError):
            report.error(row_num, f"Project_ID {project_id_raw!r} is not numeric")
            project_id = project_id_raw

        if not project:
            report.error(row_num, "Project (pathway name) is blank")

        # --- StableID -------------------------------------------------------
        if not stable_id:
            report.error(row_num, "StableID is blank")
        elif not STABLE_ID_RE.match(stable_id):
            report.error(
                row_num,
                f"StableID {stable_id!r} is malformed — expected "
                f"R-HSA-NNNNNNN.V (versioned). This string goes straight into "
                f"the deposited DOI.",
            )

        # --- StableID_root --------------------------------------------------
        derived_root = stable_id.rsplit(".", 1)[0] if "." in stable_id else stable_id
        stated_root = clean(field("stable_id_root"))
        if not stated_root:
            stable_root = derived_root
        elif stated_root != derived_root:
            report.error(
                row_num,
                f"StableID_root {stated_root!r} does not match StableID "
                f"{stable_id!r} (expected {derived_root!r}) — using the derived "
                f"value.",
            )
            stable_root = derived_root
        else:
            stable_root = stated_root

        # --- Curator --------------------------------------------------------
        cur_given, cur_surname, cur_orcid = split_curator(
            field("curator"), curator_table, row_num, report
        )

        # --- Contributor ----------------------------------------------------
        con_given, con_surname = split_contributor(
            field("contributor"), field("contributor_last"), row_num, report
        )
        if not (con_given or con_surname):
            report.error(row_num, "external contributor name is blank")

        role = clean(field("role")).title() or "Reviewer"
        if role not in {"Author", "Reviewer"}:
            report.warn(
                row_num,
                f"contributor role {role!r} is neither 'Author' nor 'Reviewer' — "
                f"passed through unchanged.",
            )

        who = f"contributor {con_given} {con_surname}".strip()
        con_orcid = resolve_orcid(field("orcid"), field("contributor_url"),
                                  row_num, who, report)

        no_orcid_flag = clean(field("no_orcid")).lower()
        if no_orcid_flag in {"yes", "y", "true"} and con_orcid:
            report.warn(
                row_num,
                f"{who} is flagged no_ORCID_ID='yes' but an ORCID is present "
                f"({con_orcid}) — using the ORCID; clear the stale flag.",
            )

        # --- Declared DOI cross-check ---------------------------------------
        declared = clean(field("declared_doi"))
        expected_doi = f"{DOI_PREFIX}/{stable_id}"
        if declared and declared != expected_doi:
            report.error(
                row_num,
                f"new-updated_DOI {declared!r} disagrees with the StableID "
                f"(expected {expected_doi!r}) — the DOI is derived from "
                f"StableID, so confirm which one is right.",
            )

        rows.append(
            {
                "Project_ID": project_id,
                "Project": project,
                "Curator Given": cur_given,
                "Curator Surname": cur_surname,
                "Curator ORCID": cur_orcid,
                "StableID": stable_id,
                "StableID_root": stable_root,
                "Contributor Given": con_given,
                "Contributor Surname": con_surname,
                "Contributor Role": role,
                "Contributor ORCID": con_orcid,
                "Editor": person_xml(cur_given, cur_surname, cur_orcid),
                "Reviewer": person_xml(con_given, con_surname, con_orcid),
                "Title": f"<title>{project}</title>",
                "DOI Data": (
                    f"<doi>{expected_doi}</doi>\n"
                    f"<resource>{RESOURCE_BASE}/{stable_id}</resource>"
                ),
                "_source_row": row_num,
            }
        )

    return pd.DataFrame(rows, columns=TARGET_COLUMNS + ["_source_row"])


def cross_row_checks(df, report):
    """Checks that only make sense across the whole sheet."""
    if df.empty:
        report.error(None, "no data rows were produced from the source sheet")
        return

    # One Project_ID must resolve to exactly one StableID — the generator groups
    # by Project_ID and takes the StableID from the first row of the group, so a
    # split here silently drops the other versions.
    for pid, group in df.groupby("Project_ID"):
        ids = sorted(set(group["StableID"]))
        if len(ids) > 1:
            rows = ", ".join(str(r) for r in group["_source_row"])
            report.error(
                None,
                f"Project_ID {pid} maps to {len(ids)} different StableIDs "
                f"({', '.join(ids)}) across source rows {rows} — the XML "
                f"generator will use only the first. Pick one.",
            )
        titles = sorted(set(group["Project"]))
        if len(titles) > 1:
            report.warn(
                None,
                f"Project_ID {pid} has {len(titles)} different titles "
                f"({'; '.join(repr(t) for t in titles)}) — the first is used.",
            )

    # The same person listed twice on the same project produces duplicate
    # contributors in the deposit.
    key = ["Project_ID", "Contributor Given", "Contributor Surname"]
    dupes = df[df.duplicated(subset=key, keep=False)]
    for (pid, given, surname), group in dupes.groupby(key):
        rows = ", ".join(str(r) for r in group["_source_row"])
        report.error(
            None,
            f"contributor {given} {surname} is listed {len(group)} times on "
            f"Project_ID {pid} (source rows {rows}) — duplicate contributors "
            f"will be deposited.",
        )

    # A StableID reused by two different projects is almost always a paste error.
    for sid, group in df.groupby("StableID"):
        pids = sorted(set(group["Project_ID"]))
        if len(pids) > 1:
            report.error(
                None,
                f"StableID {sid} is used by {len(pids)} different Project_IDs "
                f"({', '.join(str(p) for p in pids)}) — DOIs would collide.",
            )


def check_release_date(dois_path, version_number, report):
    """The generator needs a Release-sheet row for this version; check it early."""
    try:
        release = pd.read_excel(dois_path, sheet_name="Release")
    except ValueError:
        report.error(None, "DOIs.xlsx has no 'Release' sheet — "
                           "generate_crossref_xml.py will fail.")
        return None

    release.columns = release.columns.str.strip()
    if "Version" not in release.columns:
        report.error(None, "Release sheet has no 'Version' column.")
        return None

    row = release[release["Version"] == version_number]
    if row.empty:
        report.error(
            None,
            f"version {version_number} has no row in the Release sheet — add "
            f"Version/Release Year/Release Month/Release Day before running "
            f"generate_crossref_xml.py.",
        )
        return None

    row = row.iloc[0]
    return int(row["Release Year"]), int(row["Release Month"]), int(row["Release Day"])


def sort_rows(df, how):
    """Row order in prior releases: V96 groups by project, V95 sorts by surname."""
    if how == "source":
        return df
    if how == "surname":
        return df.sort_values(
            ["Contributor Surname", "Project"],
            key=lambda s: s.str.lower(),
            kind="stable",
        ).reset_index(drop=True)
    # "project" (default, V96 convention): alphabetical by pathway, contributors
    # grouped together underneath.
    return df.sort_values(
        ["Project", "Contributor Surname"],
        key=lambda s: s.str.lower(),
        kind="stable",
    ).reset_index(drop=True)


def style_sheet(ws, ncols, nrows):
    widths = [12, 52, 14, 16, 21, 18, 16, 18, 20, 16, 21, 34, 34, 34, 34]
    for i in range(1, ncols + 1):
        ws.column_dimensions[get_column_letter(i)].width = (
            widths[i - 1] if i <= len(widths) else 18
        )
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"
    for r in range(2, nrows + 2):
        for c in (12, 13, 14, 15):  # the XML helper columns
            ws.cell(row=r, column=c).alignment = Alignment(
                wrap_text=False, vertical="top"
            )


def write_standalone(df, path, sheet_name):
    path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name=sheet_name, index=False)
        style_sheet(writer.sheets[sheet_name], len(df.columns), len(df))
    return path


def write_in_place(df, dois_path, sheet_name):
    """Insert/replace the sheet inside DOIs.xlsx, after backing the file up."""
    stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    backup = dois_path.with_name(f"{dois_path.stem}.backup-{stamp}{dois_path.suffix}")
    shutil.copy2(dois_path, backup)

    wb = load_workbook(dois_path)
    replaced = sheet_name in wb.sheetnames
    if replaced:
        del wb[sheet_name]

    ws = wb.create_sheet(sheet_name)
    ws.append(list(df.columns))
    for record in df.itertuples(index=False):
        ws.append(list(record))
    style_sheet(ws, len(df.columns), len(df))

    # Keep the release sheets ahead of the lookup sheets, as in prior versions.
    anchor = next((n for n in ("Release", "Curator ORCID") if n in wb.sheetnames), None)
    if anchor:
        wb.move_sheet(sheet_name, offset=wb.sheetnames.index(anchor)
                      - wb.sheetnames.index(sheet_name))

    wb.save(dois_path)
    return backup, replaced


def main():
    parser = argparse.ArgumentParser(
        description="Format a release tracking sheet into the DOIs.xlsx layout."
    )
    parser.add_argument("worksheet", help="Target worksheet / version name, e.g. V97")
    parser.add_argument("--source", required=True,
                        help="Path to the release tracking workbook, e.g. 'Release 97.xlsx'")
    parser.add_argument("--source-sheet", default=None,
                        help="Sheet within --source (default: the worksheet name if "
                             "present, otherwise the first sheet)")
    parser.add_argument("--dois", default=None,
                        help="Path to DOIs.xlsx (for the Curator ORCID and Release "
                             "sheets, and for --in-place)")
    parser.add_argument("--output", default=None,
                        help="Write a standalone .xlsx here (default: "
                             "<worksheet>_formatted.xlsx beside --source)")
    parser.add_argument("--in-place", action="store_true",
                        help="Insert the sheet into DOIs.xlsx instead of writing a "
                             "standalone file. Backs DOIs.xlsx up first.")
    parser.add_argument("--sort", choices=["project", "surname", "source"],
                        default="project",
                        help="Row order: 'project' groups contributors under each "
                             "pathway (V96 convention, default); 'surname' matches "
                             "V95; 'source' preserves the tracking sheet's order.")
    parser.add_argument("--validate-only", action="store_true",
                        help="Report findings and write nothing.")
    parser.add_argument("--strict", action="store_true",
                        help="Exit non-zero if there are any errors.")
    args = parser.parse_args()

    source_path = Path(args.source).expanduser()
    if not source_path.exists():
        sys.exit(f"Error: source workbook not found: {source_path}")

    dois_path = Path(args.dois).expanduser() if args.dois else None
    if args.in_place and not dois_path:
        sys.exit("Error: --in-place requires --dois.")
    if dois_path and not dois_path.exists():
        sys.exit(f"Error: DOIs.xlsx not found: {dois_path}")

    version_number = int(re.sub(r"[^0-9]", "", args.worksheet))
    report = Report()

    # --- read the source sheet ---
    xl = pd.ExcelFile(source_path)
    if args.source_sheet:
        if args.source_sheet not in xl.sheet_names:
            sys.exit(f"Error: sheet {args.source_sheet!r} not found in "
                     f"{source_path.name}. Available: {xl.sheet_names}")
        source_sheet = args.source_sheet
    elif args.worksheet in xl.sheet_names:
        source_sheet = args.worksheet
    else:
        source_sheet = xl.sheet_names[0]

    src = pd.read_excel(source_path, sheet_name=source_sheet)
    src = src.loc[:, ~src.columns.astype(str).str.startswith("Unnamed:")]
    resolved = resolve_columns(src, report)

    curator_table = load_curator_orcids(dois_path, report) if dois_path else {}
    if not dois_path:
        report.warn(None, "--dois was not given — curator ORCIDs cannot be looked "
                          "up and the Release date cannot be checked.")

    df = build_rows(src, resolved, curator_table, report)
    cross_row_checks(df, report)

    release_date = check_release_date(dois_path, version_number, report) if dois_path else None

    df = sort_rows(df, args.sort)
    df = df.drop(columns=["_source_row"])

    # --- report ---
    print(f"Source:    {source_path}  (sheet {source_sheet!r})")
    print(f"Target:    worksheet {args.worksheet!r} in the DOIs.xlsx layout")
    print(f"Rows in:   {len(src)}   Rows out: {len(df)}   "
          f"Projects: {df['Project_ID'].nunique() if len(df) else 0}")
    if release_date:
        y, m, d = release_date
        print(f"Release:   {m:02d}/{d:02d}/{y}")
    report.print_report()

    if args.validate_only:
        print("\n--validate-only: nothing written.")
        return 1 if (args.strict and report.errors) else 0

    if report.errors:
        print(f"\nNOTE: {len(report.errors)} error(s) above. The sheet was still "
              f"written so you can inspect it, but resolve them before running "
              f"generate_crossref_xml.py.")

    # --- write ---
    if args.in_place:
        backup, replaced = write_in_place(df, dois_path, args.worksheet)
        verb = "Replaced" if replaced else "Added"
        print(f"\n{verb} sheet {args.worksheet!r} in {dois_path}")
        print(f"Backup:    {backup}")
    else:
        out = Path(args.output).expanduser() if args.output else \
            source_path.with_name(f"{args.worksheet}_formatted.xlsx")
        write_standalone(df, out, args.worksheet)
        print(f"\nWrote:     {out}")
        print(f"Next:      copy the {args.worksheet!r} sheet into DOIs.xlsx "
              f"(or re-run with --in-place), then run generate_crossref_xml.py.")

    return 1 if (args.strict and report.errors) else 0


if __name__ == "__main__":
    sys.exit(main())
