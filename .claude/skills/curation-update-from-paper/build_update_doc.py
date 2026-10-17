#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Render a Reactome "curation update spec" (JSON) into a formatted .docx.

Usage:
    build_update_doc.py <spec.json> [-o OUTPUT.docx]
    build_update_doc.py <spec.json> --validate

The spec describes a proposed Reactome curation update derived from a single
source paper: what the target pathway already contains, what the paper changes,
the new physical entities and events to be added, the ontology terms needed,
and the literature backing every event.

The rendered document is a *pre-curation* working document for a curator; it
does not touch the database.

Target runtime: Python 3.9+ (no match statements, no PEP 604 unions).
Requires: python-docx
"""

from __future__ import annotations

import argparse
import json
import os
import re
import shutil
import struct
import sys
import tempfile
from typing import Any, Dict, List, Optional, Sequence, Tuple

try:
    import docx  # noqa: F401
    from docx import Document
    from docx.enum.section import WD_ORIENT, WD_SECTION
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Inches, Pt, RGBColor
except ImportError:  # pragma: no cover - environment problem, not a spec problem
    sys.stderr.write(
        "error: python-docx is not installed. Install it with:\n"
        "    pip3 install --user python-docx\n"
    )
    raise SystemExit(2)

# Diagrams need matplotlib; the document is still worth producing without it,
# so this import is soft and its absence is reported once, not fatally.
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
try:
    import reaction_graphic
    GRAPHICS_AVAILABLE = True
    GRAPHICS_ERROR = ""
except (ImportError, SystemExit) as _exc:  # matplotlib missing
    reaction_graphic = None
    GRAPHICS_AVAILABLE = False
    GRAPHICS_ERROR = str(_exc) or "matplotlib not installed"

import verify_references


# ──────────────────────────────────────────────────────────────────────────────
# Constants
# ──────────────────────────────────────────────────────────────────────────────

BASE_FONT = "Calibri"
BASE_SIZE = Pt(10.5)

GREY = RGBColor(0x59, 0x59, 0x59)
LINK_BLUE = "0563C1"
HEADER_SHADE = "D9E2F3"
LABEL_SHADE = "F2F2F2"

REACTOME_DETAIL = "https://reactome.org/content/detail/"
PUBMED_BASE = "https://pubmed.ncbi.nlm.nih.gov/"

# Sections whose tables need the width. Word supports per-section orientation,
# so the rest of the document stays portrait.
LANDSCAPE_SECTIONS = frozenset(("3", "4", "5"))

# Gap analysis vocabulary. The verb leads the text in the spec ("ADDS - ...");
# it is lifted into its own colour-coded column.
GAP_VERBS = ("ADDS", "REFINES", "CONTRADICTS", "CONFIRMS")
VERB_SHADE = {
    "ADDS":        "DDEBF7",
    "REFINES":     "FFF2CC",
    "CONTRADICTS": "FBE5D6",
    "CONFIRMS":    "E2EFDA",
}
VERB_COLOR = {
    "ADDS":        RGBColor(0x1F, 0x4E, 0x79),
    "REFINES":     RGBColor(0x80, 0x60, 0x00),
    "CONTRADICTS": RGBColor(0xC0, 0x00, 0x00),
    "CONFIRMS":    RGBColor(0x37, 0x60, 0x23),
}

PRIORITY_COLOR = {
    "HIGH":   RGBColor(0xC0, 0x00, 0x00),
    "MEDIUM": RGBColor(0xBF, 0x8F, 0x00),
    "LOW":    RGBColor(0x37, 0x60, 0x23),
    "INFO":   RGBColor(0x59, 0x59, 0x59),
}
PRIORITY_SHADE = {
    "HIGH": "FBE5D6", "MEDIUM": "FFF2CC", "LOW": "E2EFDA", "INFO": "F2F2F2",
}
PRIORITY_ORDER = {"HIGH": 0, "MEDIUM": 1, "LOW": 2, "INFO": 3}

SECTIONS: Sequence[Tuple[str, str]] = (
    ("1", "Summary"),
    ("2", "Curation rationale and logic"),
    ("3", "Gap analysis against the existing pathway"),
    ("4", "Proposed new physical entities"),
    ("5", "Proposed new events"),
    ("6", "Ontology terms required"),
    ("7", "References"),
    ("8", "Open questions for the curator"),
    ("9", "Pre-submission checklist"),
)

# pmid_source vocabulary. Anything else is a validation error.
FIXED_PMID_SOURCES = frozenset(
    ("esummary:pmid", "esearch:doi", "esearch:title-author")
)
INLINE_PREFIX = "inline:"


# ──────────────────────────────────────────────────────────────────────────────
# Small defensive helpers
# ──────────────────────────────────────────────────────────────────────────────


def get(obj: Any, key: str, default: Any = None) -> Any:
    """dict.get() that tolerates None / non-dict input."""
    if isinstance(obj, dict):
        value = obj.get(key, default)
        return default if value is None else value
    return default


def raw(obj: Any, key: str) -> Any:
    """dict.get() preserving None (needed to distinguish None from False/'')."""
    if isinstance(obj, dict):
        return obj.get(key)
    return None


def as_list(value: Any) -> List[Any]:
    """Coerce a spec value into a list; tolerate a bare string or None."""
    if value is None:
        return []
    if isinstance(value, (list, tuple)):
        return [v for v in value if v is not None]
    return [value]


def as_text(value: Any) -> str:
    """Render a scalar spec value as display text."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "Yes" if value else "No"
    if isinstance(value, str):
        return value.strip()
    return str(value)


def is_blank(value: Any) -> bool:
    """True when a value should be treated as absent for rendering purposes."""
    if value is None:
        return True
    if isinstance(value, bool):
        return False
    if isinstance(value, str):
        return value.strip() == ""
    if isinstance(value, (list, tuple, dict)):
        return len(value) == 0
    return False


# ──────────────────────────────────────────────────────────────────────────────
# Low-level docx helpers (hyperlinks, shading, fields)
# ──────────────────────────────────────────────────────────────────────────────


def add_hyperlink(paragraph, url: str, text: str, color: str = LINK_BLUE,
                  underline: bool = True, size: Optional[Pt] = None,
                  italic: bool = False):
    """Append a real Word hyperlink to `paragraph`.

    python-docx has no public helper for this, so build the w:hyperlink element
    by hand and register an external relationship on the owning part.
    """
    url = (url or "").strip()
    text = text if text else url
    if not url:
        # Nothing to link to - degrade to plain text so we never lose content.
        run = paragraph.add_run(text)
        if size is not None:
            run.font.size = size
        run.italic = italic
        return None

    part = paragraph.part
    r_id = part.relate_to(url, RT.HYPERLINK, is_external=True)

    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)

    new_run = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")

    rStyle = OxmlElement("w:rStyle")
    rStyle.set(qn("w:val"), "Hyperlink")
    rPr.append(rStyle)

    if color:
        c = OxmlElement("w:color")
        c.set(qn("w:val"), color)
        rPr.append(c)
    if underline:
        u = OxmlElement("w:u")
        u.set(qn("w:val"), "single")
        rPr.append(u)
    if italic:
        i = OxmlElement("w:i")
        rPr.append(i)
    if size is not None:
        sz = OxmlElement("w:sz")
        sz.set(qn("w:val"), str(int(size.pt * 2)))
        rPr.append(sz)

    new_run.append(rPr)

    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    new_run.append(t)

    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)
    return hyperlink


def add_event_reference(paragraph, st_id: str, released: Optional[Dict[str, bool]],
                        size: Optional[Pt] = None) -> None:
    """Render an ST_ID, linking it only when the event is actually public.

    reactome.org serves released content only. A pathway still being curated
    has doRelease=false in gk_central, so a link to /content/detail/<id> 404s -
    which is exactly what made every ST_ID in the first rotavirus draft a dead
    link. When release state is known and the event is unreleased, print the
    ID as plain text with a marker instead of a link that cannot work.
    """
    st_id = (st_id or "").strip()
    if not st_id:
        return

    is_released = None if released is None else released.get(st_id)

    if is_released is False:
        run = paragraph.add_run(st_id)
        if size is not None:
            run.font.size = size
        # Covers both cases the probe cannot tell apart: an event still being
        # curated, and an ST_ID that is simply wrong.
        note = paragraph.add_run("  (not on reactome.org)")
        note.italic = True
        note.font.color.rgb = GREY
        note.font.size = Pt(8) if size is None else Pt(size.pt - 0.5)
        return

    add_hyperlink(paragraph, REACTOME_DETAIL + st_id, st_id, size=size)


def set_cell_shading(cell, hex_fill: str) -> None:
    """Apply a solid background fill to a table cell via <w:shd>."""
    tcPr = cell._tc.get_or_add_tcPr()
    for existing in tcPr.findall(qn("w:shd")):
        tcPr.remove(existing)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def repeat_header_row(row) -> None:
    """Mark a table row as a header row that repeats across page breaks."""
    trPr = row._tr.get_or_add_trPr()
    for existing in trPr.findall(qn("w:tblHeader")):
        trPr.remove(existing)
    tblHeader = OxmlElement("w:tblHeader")
    tblHeader.set(qn("w:val"), "true")
    trPr.append(tblHeader)


def set_table_autofit(table, fixed: bool = False) -> None:
    """Pick a layout algorithm: content-sized (default) or fixed.

    `fixed=True` is what makes explicit column widths stick - under the autofit
    algorithm Word recomputes them from content and ignores the widths set.
    """
    table.autofit = not fixed
    tblPr = table._tbl.tblPr
    for existing in tblPr.findall(qn("w:tblLayout")):
        tblPr.remove(existing)
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed" if fixed else "autofit")
    tblPr.append(layout)

    for existing in tblPr.findall(qn("w:tblW")):
        tblPr.remove(existing)
    tblW = OxmlElement("w:tblW")
    tblW.set(qn("w:type"), "pct")
    tblW.set(qn("w:w"), "5000")  # 100% of the text column
    tblPr.append(tblW)


def add_field(paragraph, instruction: str) -> None:
    """Insert a Word field code (e.g. PAGE, NUMPAGES) into a paragraph."""
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " %s " % instruction
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    for el in (fld_begin, instr, fld_sep, fld_end):
        run._r.append(el)


# ──────────────────────────────────────────────────────────────────────────────
# Paragraph / table conveniences
# ──────────────────────────────────────────────────────────────────────────────


def tidy_cell_paragraph(paragraph) -> None:
    """Tighten spacing inside table cells."""
    paragraph.paragraph_format.space_before = Pt(1)
    paragraph.paragraph_format.space_after = Pt(1)


def cell_paragraph(cell):
    """Return the cell's first paragraph, cleared of placeholder text."""
    paragraph = cell.paragraphs[0]
    for run in list(paragraph.runs):
        run._r.getparent().remove(run._r)
    tidy_cell_paragraph(paragraph)
    return paragraph


def write_cell_lines(cell, lines: Sequence[str], bold: bool = False,
                     size: Optional[Pt] = None) -> None:
    """Write one item per line into a cell (soft line breaks, single paragraph)."""
    paragraph = cell_paragraph(cell)
    first = True
    for line in lines:
        text = as_text(line)
        run = paragraph.add_run()
        if not first:
            run.add_break()
        run.add_text(text)
        run.bold = bold
        if size is not None:
            run.font.size = size
        first = False
    if first:  # nothing written at all
        run = paragraph.add_run("")
        if size is not None:
            run.font.size = size


def new_table(doc, headers: Sequence[str], shade: str = HEADER_SHADE):
    """Create a Table Grid table with a shaded, repeating header row."""
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    set_table_autofit(table)
    header_cells = table.rows[0].cells
    for idx, heading in enumerate(headers):
        write_cell_lines(header_cells[idx], [heading], bold=True)
        set_cell_shading(header_cells[idx], shade)
    repeat_header_row(table.rows[0])
    return table


def add_body_paragraphs(doc, paragraphs: Any) -> None:
    """Add a list of plain body paragraphs, skipping blanks."""
    items = as_list(paragraphs)
    if not items:
        add_placeholder(doc)
        return
    for item in items:
        text = as_text(item)
        if text:
            doc.add_paragraph(text)


def add_placeholder(doc, text: str = "(none recorded)") -> None:
    paragraph = doc.add_paragraph()
    run = paragraph.add_run(text)
    run.italic = True
    run.font.color.rgb = GREY


def add_bullets(doc, items: Any, prefix: Optional[str] = None) -> None:
    """Bulleted list; with `prefix` set, use plain paragraphs with that prefix."""
    values = [as_text(v) for v in as_list(items)]
    values = [v for v in values if v]
    if not values:
        add_placeholder(doc)
        return
    for value in values:
        if prefix:
            paragraph = doc.add_paragraph("%s%s" % (prefix, value))
            paragraph.paragraph_format.left_indent = Cm(0.5)
            paragraph.paragraph_format.space_after = Pt(2)
        else:
            doc.add_paragraph(value, style="List Bullet")


def add_labelled_block(doc, label: str, text: str, italic_body: bool = False,
                       quoted: bool = False) -> None:
    """A bold label paragraph followed by its body paragraph."""
    body = as_text(text)
    if not body:
        return
    label_p = doc.add_paragraph()
    label_p.paragraph_format.space_before = Pt(6)
    label_p.paragraph_format.space_after = Pt(0)
    label_run = label_p.add_run(label)
    label_run.bold = True

    if quoted and not (body.startswith("“") or body.startswith('"')):
        body = "“%s”" % body

    body_p = doc.add_paragraph()
    body_p.paragraph_format.left_indent = Cm(0.5)
    body_p.paragraph_format.space_after = Pt(4)
    body_run = body_p.add_run(body)
    body_run.italic = italic_body


# ──────────────────────────────────────────────────────────────────────────────
# Validation
# ──────────────────────────────────────────────────────────────────────────────


# Vocabulary that signals a reaction whose only change is shape. Reactome has
# no slot for conformation, so such a reaction has identical participants on
# both sides and cannot be curated as written.
_CONFORMATION_TITLE = re.compile(
    r"\bconformation(al)?\b|\btransition(s|ed)?\s+(from|to|between)\b"
    r"|\brefold|\brearrange|\bchanges?\s+shape\b|\bflips?\b"
    r"|\bupright\b|\breversed\b|\bopen(s)?\s+(up|to)\b",
    re.IGNORECASE,
)
# Qualifiers that describe a state rather than a different set of components.
_STATE_WORD = re.compile(
    r"\b(reversed|upright|open|closed|extended|compact|primed|bent|folded|"
    r"unfolded|membrane-inserted|inserted|activated|active|inactive|"
    r"conformation\w*|form|state)\b",
    re.IGNORECASE,
)


def _entity_core(name: str) -> str:
    """A name reduced to its components, with state words and bracketing gone."""
    text = re.sub(r"\([^)]*\)", " ", name or "")
    text = _STATE_WORD.sub(" ", text)
    text = re.sub(r"[^a-z0-9]+", "", text.lower())
    return text


def conformation_advisories(spec: Dict[str, Any]) -> List[str]:
    """Flag proposed events that Reactome's data model cannot represent.

    A ReactionLikeEvent records a change in the set of participants or their
    compartments. If both sides carry the same components in the same
    compartments, there is no reaction to store - however real the physical
    change. This catches the common case: a pure conformational step.
    """
    notes: List[str] = []
    for idx, event in enumerate(as_list(spec.get("new_events")), start=1):
        if not isinstance(event, dict):
            continue
        label = as_text(get(event, "label")) or "event %d" % idx
        name = as_text(get(event, "name"))

        if not GRAPHICS_AVAILABLE:
            continue
        inputs = reaction_graphic.parse_entities(raw(event, "inputs"))
        outputs = reaction_graphic.parse_entities(raw(event, "outputs"))
        if not inputs or not outputs:
            continue

        in_set = {(_entity_core(e["name"]), e["compartment"].lower())
                  for e in inputs}
        out_set = {(_entity_core(e["name"]), e["compartment"].lower())
                   for e in outputs}

        if in_set == out_set:
            notes.append(
                "%s (%s): inputs and outputs carry the same components in the "
                "same compartment, so there is no change for Reactome to "
                "store. Fold the state into the entity name of the step that "
                "does change something, or curate the binding/translocation "
                "it enables." % (label, name)
            )
        elif _CONFORMATION_TITLE.search(name) and (
            {c for _, c in in_set} == {c for _, c in out_set}
        ):
            notes.append(
                "%s (%s): reads as a conformational change within one "
                "compartment. Check that the participants really differ; "
                "Reactome cannot represent shape alone." % (label, name)
            )
    return notes


# Infection vocabulary, used only to decide whether to ask about
# relatedSpecies. Deliberately broad: the check is advisory.
_INFECTION = re.compile(
    r"\b(virus|viral|virion|bacteri\w*|pathogen\w*|infect\w*|host[- ]pathogen|"
    r"capsid|phage|parasit\w*|microb\w*)\b", re.IGNORECASE
)


def slot_advisories(spec: Dict[str, Any]) -> List[str]:
    """Check proposed events against measured slot usage.

    The numbers quoted come from data_model_slots.md, taken across released
    gk_central content. These are advisory: each rule has real exceptions.
    """
    notes: List[str] = []
    events = [e for e in as_list(spec.get("new_events")) if isinstance(e, dict)]
    if not events:
        return notes

    def label_of(event, idx):
        return as_text(get(event, "label")) or "event %d" % idx

    # 1. GO BP is a pathway-level slot: 48% of pathways, 5% of reactions.
    with_go = [label_of(e, i) for i, e in enumerate(events, 1)
               if not is_blank(raw(e, "go_bp"))]
    if with_go:
        notes.append(
            "goBiologicalProcess is proposed on %d of %d reactions (%s). It is "
            "a pathway-level slot — 48%% of released pathways carry one, but "
            "only 5%% of released reactions. Propose a GO BP term for the new "
            "subpathway instead, and drop it from the reactions unless there "
            "is a specific reason. A reaction's molecular function belongs on "
            "catalystActivity." % (len(with_go), len(events),
                                   ", ".join(with_go[:6])
                                   + (" …" if len(with_go) > 6 else ""))
        )

    # 2. Two species in one reaction have exactly two legal treatments: host in
    #    species + other in relatedSpecies, or chimeric with all species listed.
    for idx, event in enumerate(events, start=1):
        species_text = as_text(get(event, "species"))
        if not species_text:
            continue
        names = [n.strip() for n in re.split(r"[,;]| and ", species_text)
                 if n.strip()]
        if len(names) < 2:
            continue
        chimeric = as_text(get(event, "is_chimeric")).lower() in (
            "true", "yes", "1")
        if not chimeric:
            notes.append(
                "%s lists %d species in the species slot (%s) and is not marked "
                "chimeric. Two species in one reaction means either the host in "
                "species with the other in relatedSpecies, or a chimeric "
                "reaction with all species listed and isChimeric set — not a "
                "second species in the species slot."
                % (label_of(event, idx), len(names), "; ".join(names))
            )

    # 3. relatedSpecies: a branch is host-pathogen or it is not.
    with_rs = [label_of(e, i) for i, e in enumerate(events, 1)
               if not is_blank(raw(e, "related_species"))]
    without_rs = [label_of(e, i) for i, e in enumerate(events, 1)
                  if is_blank(raw(e, "related_species"))]

    if with_rs and without_rs:
        notes.append(
            "relatedSpecies is set on %d event(s) but missing from %s. A branch "
            "is either host-pathogen or it is not; make it consistent."
            % (len(with_rs), ", ".join(without_rs[:6])
               + (" …" if len(without_rs) > 6 else ""))
        )
    elif not with_rs:
        haystack = " ".join([
            as_text(spec.get("topic")), as_text(spec.get("title")),
            " ".join(as_text(p) for p in as_list(spec.get("summary"))),
        ])
        if _INFECTION.search(haystack):
            notes.append(
                "This reads as an infection pathway, and no event carries "
                "relatedSpecies. For a host-pathogen event the species slot is "
                "the host and the pathogen goes in relatedSpecies (Curator "
                "Guide pp.221-222); 82% of released RLEs carrying "
                "relatedSpecies are disease events. Check the branch in "
                "gk_central — if its events carry it, omitting it here is an "
                "incomplete annotation, not a style choice."
            )

    # 4. A value with no query behind it is not a resolved value.
    #    RESOLVED / AMBIGUOUS / JUDGEMENT all record what was checked; a bare
    #    PENDING records only that the work was skipped. See
    #    resolution_queries.md.
    text = json.dumps(spec, ensure_ascii=False)
    bare = len(re.findall(r"PENDING(?!\s*[-—–:]?\s*(RESOLVED|AMBIGUOUS))", text,
                          re.IGNORECASE))
    earned = len(re.findall(r"\b(RESOLVED|AMBIGUOUS|JUDGEMENT)\b", text))
    if bare and not earned:
        notes.append(
            "%d value(s) are marked PENDING and none is marked RESOLVED, "
            "AMBIGUOUS or JUDGEMENT. A value with no query behind it has not "
            "been checked — resolve what the database can resolve (identity, "
            "existence) and state explicitly where the open question is fit "
            "rather than identity. See resolution_queries.md." % bare
        )
    elif bare:
        notes.append(
            "%d value(s) still marked PENDING against %d resolved. Confirm each "
            "remaining one is genuinely a curator judgement rather than an "
            "unrun query." % (bare, earned)
        )

    # 5. SimpleEntity takes no species.
    for idx, ent in enumerate(as_list(spec.get("new_entities")), start=1):
        if not isinstance(ent, dict):
            continue
        if re.search(r"simple\s*entity", as_text(get(ent, "class")), re.I) \
                and not is_blank(raw(ent, "species")):
            notes.append(
                "%s is a SimpleEntity with a species. SimpleEntity carries no "
                "species in released content, at any threshold."
                % (as_text(get(ent, "proposed_name")) or "entity %d" % idx)
            )
    return notes


def collect_st_ids(spec: Dict[str, Any]) -> List[str]:
    """Every ST_ID the document will try to link, deduplicated."""
    found: List[str] = []
    pattern = re.compile(r"\b(R-[A-Z]{3}-\d+)\b")

    def walk(node: Any) -> None:
        if isinstance(node, dict):
            for value in node.values():
                walk(value)
        elif isinstance(node, (list, tuple)):
            for value in node:
                walk(value)
        elif isinstance(node, str):
            found.extend(pattern.findall(node))

    walk(spec)
    return list(dict.fromkeys(found))


def check_released(st_ids: Sequence[str], timeout: int = 15
                   ) -> Dict[str, bool]:
    """Ask reactome.org which of these ST_IDs are public.

    ContentService answers 200 for released content and 404 otherwise, which is
    exactly the released/unreleased split that decides whether a link works.
    A network failure returns {} - unknown, not "all fine".
    """
    import urllib.error
    import urllib.parse
    import urllib.request

    base = "https://reactome.org/ContentService/data/query/"
    state: Dict[str, bool] = {}
    for st_id in st_ids:
        request = urllib.request.Request(
            base + urllib.parse.quote(st_id),
            method="GET",
            headers={"User-Agent": "reactome-curation-update/1.1",
                     "Accept": "application/json"},
        )
        try:
            with urllib.request.urlopen(request, timeout=timeout) as response:
                state[st_id] = response.status == 200
        except urllib.error.HTTPError as exc:
            state[st_id] = exc.code == 200
        except (urllib.error.URLError, OSError):
            return {}   # network down: say nothing rather than something wrong
    return state


def pmid_source_ok(value: Any) -> bool:
    """None is allowed; otherwise the controlled vocabulary or inline:<something>."""
    if value is None:
        return True
    if not isinstance(value, str):
        return False
    value = value.strip()
    if value in FIXED_PMID_SOURCES:
        return True
    if value.startswith(INLINE_PREFIX) and value[len(INLINE_PREFIX):].strip():
        return True
    return False


def has_pmid(obj: Any) -> bool:
    """True when the object carries a bare PMID value."""
    value = raw(obj, "pmid")
    if value is None:
        return False
    if isinstance(value, str):
        return value.strip() != ""
    return True


def validate_spec(spec: Any) -> List[str]:
    """Return a list of human-readable validation errors (empty == valid)."""
    errors: List[str] = []

    if not isinstance(spec, dict):
        return ["spec root must be a JSON object, got %s" % type(spec).__name__]

    allowed_desc = (
        "expected one of esummary:pmid, esearch:doi, esearch:title-author, "
        "inline:<something>, or null"
    )

    # --- events: controlled vocabulary + required when a PMID is present -----
    events = as_list(spec.get("new_events"))
    for idx, event in enumerate(events):
        where = "new_events[%d]" % idx
        label = as_text(raw(event, "label")) or as_text(raw(event, "name"))
        if label:
            where = "%s (%s)" % (where, label)
        if not isinstance(event, dict):
            errors.append("%s: must be a JSON object" % where)
            continue

        source = raw(event, "pmid_source")
        if not pmid_source_ok(source):
            errors.append(
                "%s: invalid pmid_source %r - %s" % (where, source, allowed_desc)
            )
        # A verbatim quote has to be attributable, or it cannot be checked.
        if not is_blank(raw(event, "evidence_sentence")):
            citation = raw(event, "evidence_citation") or raw(
                event, "authors_year")
            if is_blank(citation):
                errors.append(
                    "%s: carries an evidence_sentence with no "
                    "evidence_citation (or authors_year). A verbatim quote "
                    "must name its source." % where
                )
            if "evidence_page" not in (event if isinstance(event, dict) else {}):
                errors.append(
                    "%s: carries an evidence_sentence but no evidence_page "
                    "key. Set it to the page the sentence was found on, or to "
                    "null if it could not be located — but record that the "
                    "question was asked." % where
                )

        if has_pmid(event) and source is None:
            errors.append(
                "%s: carries pmid %r but pmid_source is null - every bare PMID "
                "must record how it was resolved" % (where, raw(event, "pmid"))
            )

    # --- other places pmid_source appears: vocabulary check only ------------
    source_paper = spec.get("source_paper")
    if isinstance(source_paper, dict):
        source = raw(source_paper, "pmid_source")
        if not pmid_source_ok(source):
            errors.append(
                "source_paper: invalid pmid_source %r - %s" % (source, allowed_desc)
            )

    for idx, ref in enumerate(as_list(spec.get("references"))):
        if not isinstance(ref, dict):
            errors.append("references[%d]: must be a JSON object" % idx)
            continue
        source = raw(ref, "pmid_source")
        if not pmid_source_ok(source):
            errors.append(
                "references[%d]: invalid pmid_source %r - %s"
                % (idx, source, allowed_desc)
            )

    return errors


# ──────────────────────────────────────────────────────────────────────────────
# Document assembly
# ──────────────────────────────────────────────────────────────────────────────


def configure_document(doc) -> None:
    """Letter page, 2 cm margins, Calibri 10.5pt base font."""
    for section in doc.sections:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.left_margin = Cm(2)
        section.right_margin = Cm(2)
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)

    normal = doc.styles["Normal"]
    normal.font.name = BASE_FONT
    normal.font.size = BASE_SIZE
    rpr = normal.element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), BASE_FONT)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.space_before = Pt(0)

    for style_name in ("Title", "Heading 1", "Heading 2", "Heading 3",
                       "List Bullet", "List Number"):
        try:
            style = doc.styles[style_name]
        except KeyError:
            continue
        style.font.name = BASE_FONT
        srpr = style.element.get_or_add_rPr()
        srfonts = srpr.get_or_add_rFonts()
        for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            srfonts.set(qn(attr), BASE_FONT)


def add_footer_page_number(doc) -> None:
    for section in doc.sections:
        footer = section.footer
        footer.is_linked_to_previous = False
        paragraph = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
        for run in list(paragraph.runs):
            run._r.getparent().remove(run._r)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        prefix = paragraph.add_run("Page ")
        prefix.font.size = Pt(8.5)
        prefix.font.color.rgb = GREY
        add_field(paragraph, "PAGE")
        mid = paragraph.add_run(" of ")
        mid.font.size = Pt(8.5)
        mid.font.color.rgb = GREY
        add_field(paragraph, "NUMPAGES")
        for run in paragraph.runs:
            run.font.size = Pt(8.5)
            run.font.color.rgb = GREY


def add_title_block(doc, spec: Dict[str, Any]) -> None:
    title = as_text(spec.get("title")) or "Reactome curation update"
    doc.add_heading(title, level=0)

    subtitle = as_text(spec.get("subtitle"))
    if subtitle:
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(10)
        run = paragraph.add_run(subtitle)
        run.italic = True
        run.font.color.rgb = GREY
        run.font.size = Pt(11)

    pathway = spec.get("target_pathway") or {}
    paper = spec.get("source_paper") or {}

    rows: List[Tuple[str, Any, Optional[str]]] = []

    topic = as_text(spec.get("topic"))
    if topic:
        rows.append(("Topic", topic, None))

    pathway_name = as_text(get(pathway, "name"))
    pathway_st_id = as_text(get(pathway, "st_id"))
    pathway_db_id = as_text(get(pathway, "db_id"))
    pathway_url = as_text(get(pathway, "url")) or (
        REACTOME_DETAIL + pathway_st_id if pathway_st_id else ""
    )
    if pathway_name or pathway_st_id or pathway_db_id:
        bits = [b for b in (pathway_name, pathway_st_id) if b]
        label = " — ".join(bits) if bits else pathway_url
        if pathway_db_id:
            label = "%s (DB_ID %s)" % (label, pathway_db_id)
        rows.append(("Target pathway", label, pathway_url))

    citation = as_text(get(paper, "citation"))
    paper_pmid = as_text(get(paper, "pmid"))
    paper_url = as_text(get(paper, "pmid_url")) or (
        PUBMED_BASE + paper_pmid + "/" if paper_pmid else ""
    )
    paper_label = citation or ("PMID %s" % paper_pmid if paper_pmid else "")
    if paper_label:
        if paper_pmid and citation:
            paper_label = "%s  [PMID %s]" % (citation, paper_pmid)
        rows.append(("Source paper", paper_label, paper_url))

    doi = as_text(get(paper, "doi"))
    doi_url = as_text(get(paper, "doi_url")) or (
        "https://doi.org/" + doi if doi else ""
    )
    if doi:
        rows.append(("DOI", doi, doi_url))

    paper_source = raw(paper, "pmid_source")
    if not is_blank(paper_source):
        rows.append(("PMID source", as_text(paper_source), None))

    date = as_text(spec.get("date"))
    if date:
        rows.append(("Date", date, None))

    author = as_text(spec.get("author"))
    if author:
        rows.append(("Prepared by", author, None))

    rows.append((
        "References",
        "NOT verified against PubMed — do not treat citations as checked"
        if spec.get("_refs_unverified")
        else "Every PMID resolved against PubMed and matched on first author, "
             "year, journal and title",
        None,
    ))

    if not rows:
        return

    table = doc.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    set_table_autofit(table)
    for label, value, url in rows:
        cells = table.add_row().cells
        write_cell_lines(cells[0], [label], bold=True, size=Pt(9.5))
        set_cell_shading(cells[0], LABEL_SHADE)
        paragraph = cell_paragraph(cells[1])
        if url:
            add_hyperlink(paragraph, url, as_text(value), size=Pt(9.5))
        else:
            run = paragraph.add_run(as_text(value))
            run.font.size = Pt(9.5)
            if label == "References":
                unverified = bool(spec.get("_refs_unverified"))
                run.bold = unverified
                run.font.color.rgb = (
                    PRIORITY_COLOR["HIGH"] if unverified else PRIORITY_COLOR["LOW"]
                )
    doc.add_paragraph()


def heading(doc, number: str, title: str):
    return doc.add_heading("%s  %s" % (number, title), level=1)


def set_orientation(section, landscape: bool) -> None:
    """Letter, rotated or not. Word needs page_width/height swapped as well as
    the orientation flag; setting only the flag leaves a portrait text block."""
    if landscape:
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width, section.page_height = Inches(11), Inches(8.5)
    else:
        section.orientation = WD_ORIENT.PORTRAIT
        section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.left_margin = section.right_margin = Cm(1.8)
    section.top_margin = section.bottom_margin = Cm(1.8)


def content_width(doc) -> float:
    """Usable width of the current section, in cm."""
    section = doc.sections[-1]
    return (section.page_width - section.left_margin
            - section.right_margin) / 360000.0


def set_column_widths(table, widths_cm: Sequence[float]) -> None:
    """Pin column widths. Word honours this only when autofit is off and the
    width is repeated on every cell, so do both."""
    set_table_autofit(table, fixed=True)
    for row in table.rows:
        for idx, width in enumerate(widths_cm):
            if idx < len(row.cells):
                row.cells[idx].width = Cm(width)


def split_verb(text: str) -> Tuple[str, str]:
    """Separate a leading "ADDS - " style verb from the body text."""
    stripped = (text or "").lstrip()
    for verb in GAP_VERBS:
        match = re.match(
            r"^" + verb + r"\b\s*[—–:-]*\s*", stripped, re.IGNORECASE
        )
        if match:
            return verb, stripped[match.end():].lstrip()
    return "", text or ""


def render_gap_analysis(doc, rows: Any, released: Optional[Dict[str, bool]] = None
                        ) -> None:
    """Existing event (name over ST_ID) | verb | what changes.

    The old three-column layout gave a whole column to a bare ST_ID while the
    prose column did all the work. Folding the ST_ID under the event name and
    promoting the verb buys most of the width back.
    """
    items = [r for r in as_list(rows) if isinstance(r, dict)]
    if not items:
        add_placeholder(doc)
        return

    table = new_table(doc, ["Existing event", "Change", "What this paper changes"])
    for item in items:
        cells = table.add_row().cells

        name_p = cell_paragraph(cells[0])
        name_run = name_p.add_run(as_text(get(item, "existing_event")))
        name_run.bold = True
        st_id = as_text(get(item, "st_id"))
        if st_id:
            id_p = cells[0].add_paragraph()
            tidy_cell_paragraph(id_p)
            add_event_reference(id_p, st_id, released, size=Pt(8.5))

        verb, body = split_verb(as_text(get(item, "what_the_paper_changes")))
        verb_p = cell_paragraph(cells[1])
        if verb:
            verb_run = verb_p.add_run(verb)
            verb_run.bold = True
            verb_run.font.size = Pt(9)
            verb_run.font.color.rgb = VERB_COLOR.get(verb, GREY)
            set_cell_shading(cells[1], VERB_SHADE.get(verb, LABEL_SHADE))
        else:
            verb_p.add_run("—")

        write_cell_lines(cells[2], [body])

    total = content_width(doc)
    set_column_widths(table, [total * 0.26, total * 0.10, total * 0.64])


def render_new_entities(doc, rows: Any) -> None:
    items = [r for r in as_list(rows) if isinstance(r, dict)]
    if not items:
        add_placeholder(doc)
        return
    table = new_table(
        doc,
        ["Proposed name", "Class", "Compartment", "Reuses existing", "Status", "Notes"],
    )
    for item in items:
        cells = table.add_row().cells
        write_cell_lines(cells[0], [as_text(get(item, "proposed_name"))], bold=True)
        write_cell_lines(cells[1], [as_text(get(item, "class"))])
        write_cell_lines(cells[2], [as_text(get(item, "compartment"))])
        reuses = raw(item, "reuses")
        write_cell_lines(cells[3], [as_text(reuses) if not is_blank(reuses) else "—"])

        status = as_text(get(item, "status"))
        status_p = cell_paragraph(cells[4])
        status_run = status_p.add_run(status)
        status_run.font.size = Pt(9)
        # "NEW" vs "reuse" is the thing a curator scans this column for.
        if re.search(r"\bnew\b", status, re.IGNORECASE):
            status_run.bold = True
            status_run.font.color.rgb = PRIORITY_COLOR["HIGH"]
            set_cell_shading(cells[4], PRIORITY_SHADE["HIGH"])
        elif re.search(r"\breus", status, re.IGNORECASE):
            status_run.font.color.rgb = PRIORITY_COLOR["LOW"]
            set_cell_shading(cells[4], PRIORITY_SHADE["LOW"])

        write_cell_lines(cells[5], [as_text(get(item, "notes"))], size=Pt(9))

    total = content_width(doc)
    # Notes carries the prose, so give it the slack; Class/Status are short.
    set_column_widths(table, [
        total * 0.22, total * 0.10, total * 0.12,
        total * 0.14, total * 0.09, total * 0.33,
    ])


def png_size(path: str) -> Optional[Tuple[int, int]]:
    """(width, height) in pixels, read from the PNG IHDR chunk."""
    try:
        with open(path, "rb") as handle:
            header = handle.read(24)
        if len(header) < 24 or header[:8] != b"\x89PNG\r\n\x1a\n":
            return None
        return struct.unpack(">II", header[16:24])
    except OSError:
        return None


def add_scaled_picture(doc, path: str, max_w_cm: float, max_h_cm: float) -> bool:
    """Insert a picture scaled to fit inside the given box, preserving ratio.

    Without the height cap a tall reaction (several outputs plus a regulator)
    runs off the bottom of a landscape page.
    """
    size = png_size(path)
    if not size:
        return False
    width_px, height_px = size
    if not width_px or not height_px:
        return False

    width_cm = max_w_cm
    height_cm = width_cm * (height_px / float(width_px))
    if height_cm > max_h_cm:
        height_cm = max_h_cm
        width_cm = height_cm * (width_px / float(height_px))

    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.add_run().add_picture(path, width=Cm(width_cm))
    return True


def event_attribute_rows(event: Dict[str, Any], in_diagram: bool = False
                         ) -> List[Tuple[str, Any, Optional[str]]]:
    """Ordered (label, value, url) triples for one event's attribute table.

    `value` is either a string or a list of strings. Rows whose value is blank
    are dropped by the caller. With `in_diagram` set, the participant rows are
    omitted because the reaction graphic above already shows them.
    """
    pmid = as_text(get(event, "pmid"))
    pmid_url = as_text(get(event, "pmid_url")) or (
        PUBMED_BASE + pmid + "/" if pmid else ""
    )
    pmc_url = as_text(get(event, "pmc_url"))
    pmc_label = ""
    if pmc_url:
        match = re.search(r"(PMC\d+)", pmc_url, re.IGNORECASE)
        pmc_label = match.group(1).upper() if match else pmc_url

    # Participants, compartment and chain position are drawn in the diagram;
    # repeating them as table rows is what made this section 400 lines long.
    participants: List[Tuple[str, Any, Optional[str]]] = [
        ("Compartment", get(event, "compartment"), None),
        ("Inputs", as_list(raw(event, "inputs")), None),
        ("Outputs", as_list(raw(event, "outputs")), None),
        ("Catalyst", get(event, "catalyst"), None),
        ("Regulators", as_list(raw(event, "regulators")), None),
        ("Preceding event", get(event, "preceding_event"), None),
    ]
    rest: List[Tuple[str, Any, Optional[str]]] = [
        ("Type", get(event, "type"), None),
        ("Species", get(event, "species"), None),
        # Host-pathogen RLEs carry the pathogen here; see the Guide extract 3.2.
        ("relatedSpecies", get(event, "related_species"), None),
        ("GO BP", get(event, "go_bp"), None),
        ("GO MF", get(event, "go_mf"), None),
        ("Evidence type", get(event, "evidence_type"), None),
        ("Figure", get(event, "figure"), None),
        ("Authors/Year", get(event, "authors_year"), None),
        ("PMID", pmid, pmid_url),
        ("PMID source", get(event, "pmid_source"), None),
        ("PMC", pmc_label, pmc_url),
        ("Status", get(event, "status"), None),
    ]
    if in_diagram:
        return rest
    return participants[:1] + rest[:1] + participants[1:] + rest[1:]


def render_attribute_grid(doc, rows, pairs_per_row: int = 2) -> None:
    """Lay label/value pairs out in columns rather than one tall stack."""
    if not rows:
        return
    table = doc.add_table(rows=0, cols=pairs_per_row * 2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT

    for start in range(0, len(rows), pairs_per_row):
        chunk = rows[start:start + pairs_per_row]
        cells = table.add_row().cells
        for slot in range(pairs_per_row):
            label_cell = cells[slot * 2]
            value_cell = cells[slot * 2 + 1]
            if slot >= len(chunk):
                set_cell_shading(label_cell, "FFFFFF")
                continue
            lbl, value, url = chunk[slot]
            write_cell_lines(label_cell, [lbl], bold=True, size=Pt(9))
            set_cell_shading(label_cell, LABEL_SHADE)
            if url:
                paragraph = cell_paragraph(value_cell)
                add_hyperlink(paragraph, url, as_text(value), size=Pt(9))
            elif isinstance(value, (list, tuple)):
                write_cell_lines(value_cell, [as_text(v) for v in value],
                                 size=Pt(9))
            else:
                write_cell_lines(value_cell, [as_text(value)], size=Pt(9))

    total = content_width(doc)
    label_w = total * (0.13 if pairs_per_row == 2 else 0.20)
    value_w = (total - label_w * pairs_per_row) / pairs_per_row
    set_column_widths(table, [label_w, value_w] * pairs_per_row)


def add_evidence_attribution(doc, event: Dict[str, Any]) -> None:
    """Put the source under the quote: citation, page, figure, PMID.

    A verbatim sentence with no attribution cannot be checked, which is the
    whole point of quoting it. Page is omitted rather than guessed when it
    could not be located in the source.
    """
    if is_blank(raw(event, "evidence_sentence")):
        return

    citation = as_text(get(event, "evidence_citation")) or as_text(
        get(event, "authors_year"))
    page = as_text(get(event, "evidence_page"))
    figure = as_text(get(event, "figure"))
    pmid = as_text(get(event, "pmid"))
    pmid_url = as_text(get(event, "pmid_url")) or (
        PUBMED_BASE + pmid + "/" if pmid else "")

    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Cm(0.9)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(8)

    def small(text, italic=False):
        run = paragraph.add_run(text)
        run.font.size = Pt(8.5)
        run.font.color.rgb = GREY
        run.italic = italic
        return run

    small("— ")
    small(citation or "citation not recorded", italic=not citation)

    if page:
        small(", p. %s" % page)
    else:
        warn = paragraph.add_run(", page not located in the source")
        warn.font.size = Pt(8.5)
        warn.italic = True
        warn.font.color.rgb = PRIORITY_COLOR["MEDIUM"]

    if figure:
        small(", %s" % figure)
    if pmid:
        small("  PMID ")
        add_hyperlink(paragraph, pmid_url, pmid, size=Pt(8.5))


def render_event(doc, event: Dict[str, Any], index: int,
                 graphic_path: Optional[str] = None) -> None:
    label = as_text(get(event, "label"))
    name = as_text(get(event, "name"))
    if label and name:
        head = "%s — %s" % (label, name)
    else:
        head = label or name or "Event %d" % index
    doc.add_heading(head, level=2)

    embedded = False
    if graphic_path and os.path.isfile(graphic_path):
        section = doc.sections[-1]
        max_w = content_width(doc)
        max_h = ((section.page_height - section.top_margin
                  - section.bottom_margin) / 360000.0) * 0.62
        embedded = add_scaled_picture(doc, graphic_path, max_w, max_h)

    rows = [
        (lbl, val, url)
        for (lbl, val, url) in event_attribute_rows(event, in_diagram=embedded)
        if not is_blank(val)
    ]
    render_attribute_grid(doc, rows, pairs_per_row=2 if embedded else 1)

    add_labelled_block(doc, "Summation", get(event, "summation"))
    add_labelled_block(
        doc, "Evidence sentence", get(event, "evidence_sentence"),
        italic_body=True, quoted=True,
    )
    add_evidence_attribution(doc, event)

    flags = [as_text(f) for f in as_list(raw(event, "flags"))]
    flags = [f for f in flags if f]
    if flags:
        label_p = doc.add_paragraph()
        label_p.paragraph_format.space_before = Pt(6)
        label_p.paragraph_format.space_after = Pt(0)
        label_p.add_run("Flags").bold = True
        for flag in flags:
            doc.add_paragraph(flag, style="List Bullet")


def render_subpathway(doc, sub: Any, released: Optional[Dict[str, bool]]) -> None:
    """The container the new events go into.

    GO BP lives here, not on the reactions: 48% of released pathways carry one
    against 5% of reactions (see data_model_slots.md).
    """
    if not isinstance(sub, dict):
        return
    rows = [
        ("Proposed name", get(sub, "name"), None),
        ("Parent pathway", get(sub, "parent"), None),
        ("GO BP", get(sub, "go_bp"), None),
        ("Species", get(sub, "species"), None),
        ("relatedSpecies", get(sub, "related_species"), None),
        ("Events", get(sub, "events"), None),
    ]
    rows = [(l, v, u) for (l, v, u) in rows if not is_blank(v)]
    if not rows:
        return

    heading_p = doc.add_paragraph()
    heading_p.paragraph_format.space_after = Pt(2)
    run = heading_p.add_run("Proposed containing subpathway")
    run.bold = True
    render_attribute_grid(doc, rows, pairs_per_row=2)

    summation = get(sub, "summation")
    if not is_blank(summation):
        add_labelled_block(doc, "Summation", summation)
    doc.add_paragraph()


def render_new_events(doc, events: Any,
                      graphics: Optional[Dict[str, str]] = None) -> None:
    items = [e for e in as_list(events) if isinstance(e, dict)]
    if not items:
        add_placeholder(doc)
        return
    graphics = graphics or {}
    for idx, event in enumerate(items, start=1):
        label = as_text(get(event, "label"))
        render_event(doc, event, idx, graphic_path=graphics.get(label))


_PRIORITY_HINTS = (
    (re.compile(r"highest priority|structural decision|contradiction",
                re.IGNORECASE), "HIGH"),
    (re.compile(r"\bcorrected\b|species convention|needs? a decision|"
                r"warrant merging|cannot serve", re.IGNORECASE), "MEDIUM"),
    (re.compile(r"pending curator verification|no ehld change|"
                r"was not performed|may reasonably be declined",
                re.IGNORECASE), "LOW"),
)


def normalise_question(item: Any, index: int) -> Dict[str, str]:
    """Accept either a structured question or the older plain string.

    Structured form (preferred):
        {"priority": "HIGH", "tag": "STRUCTURAL DECISION", "question": "...",
         "affects": "E6-E8"}
    A plain string is parsed for a leading ALL-CAPS tag and scored against the
    hints above, so older specs keep rendering.
    """
    if isinstance(item, dict):
        text = as_text(get(item, "question")) or as_text(get(item, "text"))
        priority = (as_text(get(item, "priority")) or "").upper()
        tag = as_text(get(item, "tag"))
        affects = as_text(get(item, "affects"))
    else:
        text = as_text(item)
        priority, tag, affects = "", "", ""

    if not tag:
        match = re.match(r"^\s*([A-Z][A-Z /_-]{3,40}?)\s*[:,]\s*", text)
        if match:
            tag = match.group(1).strip()
            text = text[match.end():].strip()
    elif text.upper().startswith(tag.upper()):
        text = text[len(tag):].lstrip(" :,-")

    if priority not in PRIORITY_COLOR:
        priority = ""
        haystack = "%s %s" % (tag, text)
        for pattern, level in _PRIORITY_HINTS:
            if pattern.search(haystack):
                priority = level
                break
        priority = priority or "MEDIUM"

    return {"priority": priority, "tag": tag, "question": text,
            "affects": affects, "n": str(index)}


def render_open_questions(doc, items: Any,
                          flow_image: Optional[str] = None) -> None:
    """A prioritised, tickable table, optionally under a decision flowchart."""
    questions = [normalise_question(q, i)
                 for i, q in enumerate(as_list(items), start=1)]
    questions = [q for q in questions if q["question"]]
    if not questions:
        add_placeholder(doc)
        return

    if flow_image and os.path.isfile(flow_image):
        section = doc.sections[-1]
        max_h = ((section.page_height - section.top_margin
                  - section.bottom_margin) / 360000.0) * 0.42
        if add_scaled_picture(doc, flow_image, content_width(doc), max_h):
            caption = doc.add_paragraph()
            caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cap_run = caption.add_run(
                "Decision flow — the order these questions need answering in."
            )
            cap_run.italic = True
            cap_run.font.size = Pt(8.5)
            cap_run.font.color.rgb = GREY

    questions.sort(key=lambda q: PRIORITY_ORDER.get(q["priority"], 9))

    table = new_table(doc, ["", "Priority", "Question", "Affects"])
    for item in questions:
        cells = table.add_row().cells

        box_p = cell_paragraph(cells[0])
        box_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        box_p.add_run("☐").font.size = Pt(12)

        priority = item["priority"]
        pr_p = cell_paragraph(cells[1])
        pr_run = pr_p.add_run(priority)
        pr_run.bold = True
        pr_run.font.size = Pt(9)
        pr_run.font.color.rgb = PRIORITY_COLOR.get(priority, GREY)
        set_cell_shading(cells[1], PRIORITY_SHADE.get(priority, LABEL_SHADE))

        q_cell = cells[2]
        q_p = cell_paragraph(q_cell)
        if item["tag"]:
            tag_run = q_p.add_run(item["tag"] + "  ")
            tag_run.bold = True
            tag_run.font.size = Pt(9)
            tag_run.font.color.rgb = PRIORITY_COLOR.get(priority, GREY)
        q_p.add_run(item["question"])

        write_cell_lines(cells[3], [item["affects"] or "—"], size=Pt(9))

    total = content_width(doc)
    set_column_widths(table, [total * 0.05, total * 0.11,
                              total * 0.70, total * 0.14])


def render_ontology_terms(doc, rows: Any) -> None:
    items = [r for r in as_list(rows) if isinstance(r, dict)]
    if not items:
        add_placeholder(doc)
        return
    table = new_table(doc, ["Term label", "Ontology", "Accession / status", "Used by"])
    for item in items:
        cells = table.add_row().cells
        write_cell_lines(cells[0], [as_text(get(item, "term_label"))])
        write_cell_lines(cells[1], [as_text(get(item, "ontology"))])
        write_cell_lines(cells[2], [as_text(get(item, "accession_status"))])
        used_by = raw(item, "used_by")
        if isinstance(used_by, (list, tuple)):
            write_cell_lines(cells[3], [as_text(v) for v in used_by])
        else:
            write_cell_lines(cells[3], [as_text(used_by)])


def render_references(doc, refs: Any) -> None:
    items = [r for r in as_list(refs) if isinstance(r, dict)]
    if not items:
        add_placeholder(doc)
        return
    for idx, ref in enumerate(items, start=1):
        number = raw(ref, "n")
        number_text = as_text(number) if not is_blank(number) else str(idx)

        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.left_indent = Cm(0.9)
        paragraph.paragraph_format.first_line_indent = Cm(-0.9)
        paragraph.paragraph_format.space_after = Pt(2)

        num_run = paragraph.add_run("%s. " % number_text)
        num_run.bold = True
        paragraph.add_run(as_text(get(ref, "citation")))

        pmid = as_text(get(ref, "pmid"))
        pmid_url = as_text(get(ref, "pmid_url")) or (
            PUBMED_BASE + pmid + "/" if pmid else ""
        )
        if pmid:
            paragraph.add_run("  PMID ")
            add_hyperlink(paragraph, pmid_url, pmid)
        elif pmid_url:
            paragraph.add_run("  ")
            add_hyperlink(paragraph, pmid_url, pmid_url)

        note = as_text(get(ref, "note"))
        if note:
            note_run = paragraph.add_run("  — %s" % note)
            note_run.font.size = Pt(9.5)

        source = raw(ref, "pmid_source")
        if not is_blank(source):
            meta = doc.add_paragraph()
            meta.paragraph_format.left_indent = Cm(0.9)
            meta.paragraph_format.space_after = Pt(6)
            meta_run = meta.add_run("PMID source: %s" % as_text(source))
            meta_run.italic = True
            meta_run.font.size = Pt(8.5)
            meta_run.font.color.rgb = GREY


def build_document(spec: Dict[str, Any], output_path: str,
                   graphics: Optional[Dict[str, str]] = None,
                   flow_image: Optional[str] = None,
                   released: Optional[Dict[str, bool]] = None) -> str:
    doc = Document()
    configure_document(doc)
    add_footer_page_number(doc)
    add_title_block(doc, spec)

    landscape_now = False
    for number, title in SECTIONS:
        want_landscape = number in LANDSCAPE_SECTIONS
        if want_landscape != landscape_now:
            # Orientation is a section property, so each flip needs a break.
            section = doc.add_section(WD_SECTION.NEW_PAGE)
            set_orientation(section, want_landscape)
            landscape_now = want_landscape

        heading(doc, number, title)
        if number == "1":
            add_body_paragraphs(doc, spec.get("summary"))
        elif number == "2":
            add_body_paragraphs(doc, spec.get("rationale"))
        elif number == "3":
            render_gap_analysis(doc, spec.get("gap_analysis"), released)
        elif number == "4":
            render_new_entities(doc, spec.get("new_entities"))
        elif number == "5":
            render_subpathway(doc, spec.get("proposed_subpathway"), released)
            render_new_events(doc, spec.get("new_events"), graphics)
        elif number == "6":
            render_ontology_terms(doc, spec.get("ontology_terms"))
        elif number == "7":
            render_references(doc, spec.get("references"))
        elif number == "8":
            render_open_questions(doc, spec.get("open_questions"), flow_image)
        elif number == "9":
            add_bullets(doc, spec.get("checklist"), prefix="☐  ")

    doc.save(output_path)
    return output_path


# ──────────────────────────────────────────────────────────────────────────────
# CLI
# ──────────────────────────────────────────────────────────────────────────────


def default_output_path(spec: Dict[str, Any], spec_path: str) -> str:
    basename = as_text(spec.get("output_basename"))
    if not basename:
        basename = os.path.splitext(os.path.basename(spec_path))[0]
    if basename.lower().endswith(".docx"):
        basename = basename[: -len(".docx")]
    directory = os.path.dirname(os.path.abspath(spec_path))
    return os.path.join(directory, basename + ".docx")


def main(argv: Optional[Sequence[str]] = None) -> int:
    parser = argparse.ArgumentParser(
        description="Render a Reactome curation update spec (JSON) into a .docx."
    )
    parser.add_argument("spec", help="path to the curation update spec JSON file")
    parser.add_argument(
        "-o", "--output", default=None,
        help="output .docx path (default: spec's output_basename + .docx, "
             "written beside the spec file)",
    )
    parser.add_argument(
        "--validate", action="store_true",
        help="validate the spec and exit without writing a document",
    )
    parser.add_argument(
        "--skip-ref-check", action="store_true",
        help="do not resolve references against PubMed. The document is then "
             "stamped as unverified - use only when offline.",
    )
    parser.add_argument(
        "--skip-link-check", action="store_true",
        help="do not ask reactome.org which ST_IDs are released. Unreleased "
             "IDs will then be rendered as links that 404.",
    )
    parser.add_argument(
        "--no-graphics", action="store_true",
        help="skip the reaction diagrams and render the full attribute tables",
    )
    parser.add_argument(
        "--keep-figures", metavar="DIR", default=None,
        help="also write the generated PNGs to DIR (default: a temp dir, "
             "removed after the document is written)",
    )
    args = parser.parse_args(argv)

    spec_path = os.path.abspath(os.path.expanduser(args.spec))
    if not os.path.isfile(spec_path):
        sys.stderr.write("error: spec file not found: %s\n" % spec_path)
        return 1

    try:
        with open(spec_path, "r", encoding="utf-8") as handle:
            spec = json.load(handle)
    except ValueError as exc:
        sys.stderr.write("error: %s is not valid JSON: %s\n" % (spec_path, exc))
        return 1
    except OSError as exc:
        sys.stderr.write("error: could not read %s: %s\n" % (spec_path, exc))
        return 1

    errors = validate_spec(spec)
    if errors:
        sys.stderr.write("Spec validation failed (%d problem%s):\n"
                         % (len(errors), "" if len(errors) == 1 else "s"))
        for error in errors:
            sys.stderr.write("  - %s\n" % error)
        return 1

    if args.validate:
        print("Spec OK: %s" % spec_path)
        return 0

    output_path = (
        os.path.abspath(os.path.expanduser(args.output))
        if args.output else default_output_path(spec, spec_path)
    )
    out_dir = os.path.dirname(output_path)
    if out_dir and not os.path.isdir(out_dir):
        sys.stderr.write("error: output directory does not exist: %s\n" % out_dir)
        return 1

    # ── references: a wrong PMID must stop the build, not be discovered later
    if args.skip_ref_check:
        spec = dict(spec)
        spec["_refs_unverified"] = True
        sys.stderr.write(
            "warning: --skip-ref-check set; references were NOT resolved "
            "against PubMed and the document says so.\n"
        )
    else:
        try:
            result = verify_references.verify_spec(spec)
        except Exception as exc:  # network, parse, anything
            sys.stderr.write(
                "error: could not verify references against PubMed: %s\n"
                "Nothing was verified. Re-run with network access, or pass "
                "--skip-ref-check to render an explicitly unverified draft.\n"
                % exc
            )
            return 1
        print(verify_references.format_report(result))
        if result["errors"]:
            sys.stderr.write(
                "\nerror: %d reference(s) do not match the PMID recorded for "
                "them. Refusing to render.\nFix the spec, or pass "
                "--skip-ref-check if you know better than PubMed.\n"
                % len(result["errors"])
            )
            return 1

    # ── events the data model cannot hold, and slots used against practice ──
    advisories = conformation_advisories(spec) + slot_advisories(spec)
    if advisories:
        sys.stderr.write(
            "\nData model check — %d finding%s:\n"
            % (len(advisories), "" if len(advisories) == 1 else "s")
        )
        for note in advisories:
            sys.stderr.write("  - %s\n" % note)
        sys.stderr.write(
            "These are advisory, not blocking. Resolve them before creating "
            "instances.\n\n"
        )

    # ── ST_ID links: only link what reactome.org actually serves ────────────
    released: Optional[Dict[str, bool]] = None
    if not args.skip_link_check:
        st_ids = collect_st_ids(spec)
        if st_ids:
            released = check_released(st_ids) or None
            if released is None:
                sys.stderr.write(
                    "warning: could not reach reactome.org; ST_IDs will be "
                    "linked without checking whether they are released.\n"
                )
            else:
                unreleased = sorted(k for k, v in released.items() if not v)
                if unreleased:
                    print("Unreleased ST_IDs (rendered as plain text, not "
                          "links): %s" % ", ".join(unreleased))

    # ── figures ─────────────────────────────────────────────────────────────
    graphics: Dict[str, str] = {}
    flow_image: Optional[str] = None
    fig_dir = args.keep_figures
    temp_dir: Optional[str] = None
    if not args.no_graphics:
        if not GRAPHICS_AVAILABLE:
            sys.stderr.write(
                "warning: reaction diagrams skipped (%s). Install matplotlib "
                "to enable them.\n" % GRAPHICS_ERROR
            )
        else:
            if fig_dir:
                fig_dir = os.path.abspath(os.path.expanduser(fig_dir))
                os.makedirs(fig_dir, exist_ok=True)
            else:
                temp_dir = tempfile.mkdtemp(prefix="curation-update-figs-")
                fig_dir = temp_dir
            try:
                graphics = reaction_graphic.render_all(spec, fig_dir)
                flow = spec.get("decision_flow")
                if flow:
                    flow_image = reaction_graphic.render_decision_flow(
                        flow, os.path.join(fig_dir, "decision_flow.png")
                    )
            except Exception as exc:
                sys.stderr.write(
                    "warning: diagram rendering failed (%s); falling back to "
                    "tables.\n" % exc
                )
                graphics = {}

    try:
        build_document(spec, output_path, graphics=graphics,
                       flow_image=flow_image, released=released)
    except OSError as exc:
        sys.stderr.write("error: could not write %s: %s\n" % (output_path, exc))
        return 1
    finally:
        if temp_dir:
            shutil.rmtree(temp_dir, ignore_errors=True)

    print("Wrote %s" % output_path)
    if graphics:
        print("Embedded %d reaction diagram%s%s"
              % (len(graphics), "" if len(graphics) == 1 else "s",
                 " + decision flow" if flow_image else ""))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
